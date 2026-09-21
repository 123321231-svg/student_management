import io
import os
import secrets
from pathlib import Path

from fastapi import FastAPI, File, Form, Query, Request, UploadFile
from fastapi.responses import HTMLResponse, RedirectResponse, StreamingResponse
from fastapi.templating import Jinja2Templates
from openpyxl import Workbook, load_workbook
from sqlalchemy import select
from starlette.middleware.sessions import SessionMiddleware

from .api_v1 import create_api_router
from .db import connection, init_db, write_audit
from .orm import Assessment, ClassGroup, Course, ScoreRecord, Student, init_orm, orm_session, seed_admin
from .security import hash_password, new_csrf_token, verify_password
from .services import build_score_stats, parse_student, student_export_values

BASE_DIR = Path(__file__).resolve().parent
VERSION = "3.0.0"
templates = Jinja2Templates(directory=str(BASE_DIR / "templates"))
app = FastAPI(title="学生学业数据分析与管理平台", version=VERSION)
app.add_middleware(
    SessionMiddleware,
    secret_key=os.getenv("SECRET_KEY", "change-this-secret-key-in-production"),
    https_only=os.getenv("COOKIE_SECURE", "0") == "1",
    same_site="lax",
)


def startup() -> None:
    init_db()
    init_orm()
    seed_admin()


startup()


def csrf(request: Request) -> str:
    token = request.session.get("csrf_token")
    if not token:
        token = new_csrf_token()
        request.session["csrf_token"] = token
    return token


def valid_csrf(request: Request, token: str) -> bool:
    expected = request.session.get("csrf_token", "")
    return bool(token) and secrets.compare_digest(token, expected)


def current_user(request: Request):
    user_id = request.session.get("user_id")
    if not user_id:
        return None
    with connection() as conn:
        return conn.execute(
            "SELECT id, username, full_name, role, is_active FROM users WHERE id = ?",
            (user_id,),
        ).fetchone()


def render(request: Request, template: str, **context):
    user = current_user(request)
    context.update(
        user=user,
        csrf_token=csrf(request),
        message=request.session.pop("message", None),
    )
    return templates.TemplateResponse(request=request, name=template, context=context)


def redirect_login(message: str | None = None):
    response = RedirectResponse("/login", status_code=303)
    if message:
        response.headers["X-Login-Message"] = message
    return response


def require_user(request: Request):
    user = current_user(request)
    if not user or not user["is_active"]:
        return None
    return user


def require_write_user(request: Request):
    user = require_user(request)
    if not user or user["role"] not in {"admin", "teacher"}:
        return None
    return user


app.include_router(create_api_router(require_user))


@app.get("/", response_class=HTMLResponse)
def index(request: Request):
    return RedirectResponse("/dashboard" if require_user(request) else "/login", status_code=303)


@app.get("/health")
def health():
    return {"status": "ok", "service": "student-management", "version": VERSION}


@app.get("/login", response_class=HTMLResponse)
def login_page(request: Request):
    return render(request, "login.html", title="登录")


@app.post("/login")
def login(
    request: Request,
    username: str = Form(...),
    password: str = Form(...),
    csrf_token: str = Form(...),
):
    if not valid_csrf(request, csrf_token):
        return render(request, "login.html", title="登录", error="页面已过期，请刷新后重试")
    with connection() as conn:
        user = conn.execute("SELECT * FROM users WHERE username = ?", (username.strip(),)).fetchone()
        if not user or not user["is_active"] or not verify_password(password, user["password_hash"]):
            return render(request, "login.html", title="登录", error="用户名或密码错误")
        request.session["user_id"] = user["id"]
        write_audit(conn, user["id"], "login", "用户登录")
    return RedirectResponse("/dashboard", status_code=303)


@app.get("/register", response_class=HTMLResponse)
def register_page(request: Request):
    return render(request, "register.html", title="注册")


@app.post("/register")
def register(
    request: Request,
    username: str = Form(...),
    full_name: str = Form(...),
    password: str = Form(...),
    confirm_password: str = Form(...),
    csrf_token: str = Form(...),
):
    if not valid_csrf(request, csrf_token):
        return render(request, "register.html", title="注册", error="页面已过期，请刷新后重试")
    username, full_name = username.strip(), full_name.strip()
    if len(username) < 3 or len(password) < 6 or not full_name:
        return render(request, "register.html", title="注册", error="用户名至少3位，密码至少6位，姓名不能为空")
    if password != confirm_password:
        return render(request, "register.html", title="注册", error="两次输入的密码不一致")
    with connection() as conn:
        try:
            conn.execute(
                "INSERT INTO users(username, full_name, password_hash, role) VALUES (?, ?, ?, 'viewer')",
                (username, full_name, hash_password(password)),
            )
            created_user = conn.execute("SELECT id FROM users WHERE username = ?", (username,)).fetchone()
            write_audit(conn, created_user["id"], "register", "用户注册")
        except Exception:
            return render(request, "register.html", title="注册", error="用户名已存在")
    request.session["message"] = "注册成功，请登录"
    return RedirectResponse("/login", status_code=303)


@app.get("/logout")
def logout(request: Request):
    request.session.clear()
    return RedirectResponse("/login", status_code=303)


@app.get("/dashboard", response_class=HTMLResponse)
def dashboard(request: Request):
    if not require_user(request):
        return redirect_login()
    return render(request, "dashboard.html", title="仪表盘")


@app.get("/academics", response_class=HTMLResponse)
def academics(request: Request):
    if not require_user(request):
        return redirect_login()
    with orm_session() as session:
        classes = session.scalars(select(ClassGroup).order_by(ClassGroup.grade.desc(), ClassGroup.name)).all()
        courses = session.scalars(select(Course).order_by(Course.code)).all()
        assessments = session.scalars(select(Assessment).order_by(Assessment.exam_date.desc())).all()
        class_rows = [{"id": row.id, "name": row.name, "grade": row.grade, "major": row.major} for row in classes]
        course_rows = [
            {"id": row.id, "code": row.code, "name": row.name, "credits": row.credits, "teacher_name": row.teacher_name}
            for row in courses
        ]
        assessment_rows = [
            {
                "id": row.id,
                "name": row.name,
                "course_name": row.course.name,
                "exam_date": row.exam_date,
                "max_score": row.max_score,
            }
            for row in assessments
        ]
    return render(
        request,
        "academics.html",
        title="教务数据",
        classes=class_rows,
        courses=course_rows,
        assessments=assessment_rows,
    )


@app.get("/analytics", response_class=HTMLResponse)
def analytics_page(request: Request):
    if not require_user(request):
        return redirect_login()
    with orm_session() as session:
        students = session.scalars(select(Student)).all()
        scores = session.scalars(select(ScoreRecord)).all()
        failed_counts: dict[int, int] = {}
        for score in scores:
            if score.score < 60:
                failed_counts[score.student_id] = failed_counts.get(score.student_id, 0) + 1
        risk_ids = {student_id for student_id, count in failed_counts.items() if count >= 2}
        risk_students = [
            {"id": student.id, "student_id": student.student_id, "name": student.name, "failed_count": failed_counts[student.id]}
            for student in students
            if student.id in risk_ids
        ]
    return render(request, "analytics.html", title="数据分析中心", risk_students=risk_students)


@app.get("/students", response_class=HTMLResponse)
def students(request: Request, q: str = Query(default="")):
    if not require_user(request):
        return redirect_login()
    q = q.strip()
    with connection() as conn:
        rows = conn.execute(
            """
            SELECT id, student_id, name, age, score
            FROM students
            WHERE (? = '' OR student_id LIKE ? OR name LIKE ?)
            ORDER BY id DESC
            """,
            (q, f"%{q}%", f"%{q}%"),
        ).fetchall()
    return render(request, "students.html", title="学生管理", students=rows, query=q)


@app.post("/students/create")
def create_student(
    request: Request,
    student_id: str = Form(...),
    name: str = Form(...),
    age: str = Form(...),
    score: str = Form(...),
    csrf_token: str = Form(...),
):
    user = require_write_user(request)
    if not user:
        return RedirectResponse("/students", status_code=303)
    if not valid_csrf(request, csrf_token):
        request.session["message"] = "页面已过期，请刷新后重试"
        return RedirectResponse("/students", status_code=303)
    try:
        values = parse_student(student_id, name, age, score)
        with connection() as conn:
            conn.execute("INSERT INTO students(student_id, name, age, score) VALUES (?, ?, ?, ?)", values)
            write_audit(conn, user["id"], "create_student", f"创建学生 {values[0]}")
        request.session["message"] = "学生添加成功"
    except ValueError as error:
        request.session["message"] = str(error)
    except Exception:
        request.session["message"] = "添加失败，学号可能已经存在"
    return RedirectResponse("/students", status_code=303)


@app.post("/students/{student_id}/update")
def update_student(
    student_id: int,
    request: Request,
    name: str = Form(...),
    age: str = Form(...),
    score: str = Form(...),
    csrf_token: str = Form(...),
):
    user = require_write_user(request)
    if not user:
        return RedirectResponse("/students", status_code=303)
    try:
        values = parse_student("placeholder", name, age, score)
        with connection() as conn:
            conn.execute("UPDATE students SET name = ?, age = ?, score = ? WHERE id = ?", values[1:] + (student_id,))
            write_audit(conn, user["id"], "update_student", f"更新学生记录 {student_id}")
        request.session["message"] = "学生修改成功"
    except ValueError as error:
        request.session["message"] = str(error)
    except Exception:
        request.session["message"] = "学生修改失败"
    return RedirectResponse("/students", status_code=303)


@app.post("/students/{student_id}/delete")
def delete_student(student_id: int, request: Request, csrf_token: str = Form(...)):
    user = require_write_user(request)
    if user and valid_csrf(request, csrf_token):
        with connection() as conn:
            conn.execute("DELETE FROM students WHERE id = ?", (student_id,))
            write_audit(conn, user["id"], "delete_student", f"删除学生记录 {student_id}")
        request.session["message"] = "学生删除成功"
    return RedirectResponse("/students", status_code=303)


@app.get("/students/export")
def export_students(request: Request, q: str = Query(default="")):
    if not require_user(request):
        return redirect_login()
    q = q.strip()
    with connection() as conn:
        rows = conn.execute(
            "SELECT student_id, name, age, score FROM students WHERE (? = '' OR student_id LIKE ? OR name LIKE ?) ORDER BY student_id",
            (q, f"%{q}%", f"%{q}%"),
        ).fetchall()
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "学生信息"
    sheet.append(["学号", "姓名", "年龄", "成绩"])
    for row in rows:
        sheet.append(student_export_values(row))
    for column in sheet.columns:
        sheet.column_dimensions[column[0].column_letter].width = 18
    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=students.xlsx"},
    )


@app.post("/students/import")
async def import_students(request: Request, file: UploadFile = File(...), csrf_token: str = Form(...)):
    user = require_write_user(request)
    if not user or not valid_csrf(request, csrf_token):
        return RedirectResponse("/students", status_code=303)
    try:
        workbook = load_workbook(io.BytesIO(await file.read()), read_only=True, data_only=True)
        sheet = workbook.active
        headers = [str(cell.value or "").strip() for cell in next(sheet.iter_rows(max_row=1))]
        expected = ["学号", "姓名", "年龄", "成绩"]
        if headers[:4] != expected:
            raise ValueError("Excel第一行必须依次为：学号、姓名、年龄、成绩")
        imported = 0
        with connection() as conn:
            for row in sheet.iter_rows(min_row=2, values_only=True):
                if not any(value is not None and str(value).strip() for value in row[:4]):
                    continue
                values = parse_student(str(row[0] or ""), str(row[1] or ""), str(row[2] or ""), str(row[3] or ""))
                conn.execute(
                    "INSERT INTO students(student_id, name, age, score) VALUES (?, ?, ?, ?) ON CONFLICT(student_id) DO UPDATE SET name=excluded.name, age=excluded.age, score=excluded.score",
                    values,
                )
                imported += 1
            write_audit(conn, user["id"], "import_students", f"Excel导入 {imported} 条")
        request.session["message"] = f"成功导入 {imported} 条学生记录"
    except Exception as error:
        request.session["message"] = f"导入失败：{error}"
    return RedirectResponse("/students", status_code=303)


@app.get("/api/stats")
def stats(request: Request):
    if not require_user(request):
        return {"error": "未登录"}
    with connection() as conn:
        rows = conn.execute("SELECT score FROM students").fetchall()
    return build_score_stats([float(row["score"]) for row in rows])


@app.get("/users", response_class=HTMLResponse)
def users(request: Request):
    user = require_user(request)
    if not user or user["role"] != "admin":
        return RedirectResponse("/dashboard", status_code=303)
    with connection() as conn:
        rows = conn.execute("SELECT id, username, full_name, role, is_active, created_at FROM users ORDER BY id").fetchall()
    return render(request, "users.html", title="用户与权限", users=rows)


@app.get("/audit-logs", response_class=HTMLResponse)
def audit_logs(request: Request):
    user = require_user(request)
    if not user or user["role"] != "admin":
        return RedirectResponse("/dashboard", status_code=303)
    with connection() as conn:
        rows = conn.execute(
            """
            SELECT audit_logs.id, audit_logs.action, audit_logs.detail,
                   audit_logs.created_at, users.full_name
            FROM audit_logs
            LEFT JOIN users ON users.id = audit_logs.user_id
            ORDER BY audit_logs.id DESC
            LIMIT 200
            """
        ).fetchall()
    return render(request, "audit_logs.html", title="审计日志", logs=rows)


@app.post("/users/{user_id}/role")
def update_role(user_id: int, request: Request, role: str = Form(...), csrf_token: str = Form(...)):
    user = require_user(request)
    if not user or user["role"] != "admin" or not valid_csrf(request, csrf_token):
        return RedirectResponse("/dashboard", status_code=303)
    if role not in {"admin", "teacher", "viewer"}:
        role = "viewer"
    with connection() as conn:
        target = conn.execute("SELECT role FROM users WHERE id = ?", (user_id,)).fetchone()
        admin_count = conn.execute(
            "SELECT COUNT(*) AS count FROM users WHERE role = 'admin' AND is_active = 1"
        ).fetchone()["count"]
        if target and target["role"] == "admin" and role != "admin" and admin_count <= 1:
            request.session["message"] = "系统至少需要保留一名管理员"
            return RedirectResponse("/users", status_code=303)
        conn.execute("UPDATE users SET role = ? WHERE id = ?", (role, user_id))
        write_audit(conn, user["id"], "update_role", f"用户 {user_id} 设置为 {role}")
    request.session["message"] = "权限更新成功"
    return RedirectResponse("/users", status_code=303)
